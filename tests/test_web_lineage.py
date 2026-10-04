"""Lineage tab (L2): pages, filters, JSON API, CSV / Excel exports, auth, escaping, row cap."""

import asyncio
import base64
import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from pch.demo.generator import generate_world
from pch.demo.transport import parse_world_time
from pch.orchestrator import ScanConfig, Scanner
from pch.settings import Policy, Scope, Settings
from pch.sources import demo_sources
from pch.store.db import session_scope
from pch.store.models import LineageRow
from pch.web import exports as X
from pch.web import lineage_q as LQ
from pch.web.app import create_app
from tests.builders import ALL_HOSTS

ROLE = "PCH.Reader"
PAYLOADS = ["<img src=x onerror=alert(1)>", "</script><script>alert(1)</script>", '"><svg/onload=alert(1)>', "javascript:alert(1)"]
FORMULAS = ["=HYPERLINK(\"http://evil\",\"x\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)"]


def S(**kw) -> Settings:
    return Settings(_env_file=None, **kw)  # type: ignore[call-arg]


def make_db(tmp: Path, repos=120, seed=7) -> str:
    url = f"sqlite:///{tmp}/lin.db"
    w = generate_world(seed=seed, repos=repos, now=datetime(2026, 10, 1, 12, 0, 0))
    src = demo_sources(w)
    cfg = ScanConfig(scope=Scope(code_hosts=ALL_HOSTS, projects=w["meta"]["projects"]), policy=Policy(approved_registries=["contosoacr.azurecr.io"]), db_url=url, mode="demo", now=parse_world_time(w))

    async def go():
        try:
            await Scanner(src, cfg).run("scan-0")
        finally:
            await src.aclose()

    asyncio.run(go())
    return url


@pytest.fixture(scope="module")
def db(tmp_path_factory):
    return make_db(tmp_path_factory.mktemp("lin"))


@pytest.fixture(scope="module")
def client(db):
    return TestClient(create_app(db, settings=S()), base_url="http://localhost")


def copy_db(db: str, dest: Path) -> None:
    """Consistent copy of a SQLite file (the source may be in WAL mode: a plain file copy would miss recent commits)."""
    import sqlite3

    src = sqlite3.connect(db.removeprefix("sqlite:///"))
    out = sqlite3.connect(dest)
    src.backup(out)
    out.close()
    src.close()


def rows_of(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


# ------------------------------------------------------------------ pages
def test_lineage_is_the_last_nav_item_and_page_renders(client):
    r = client.get("/lineage")
    assert r.status_code == 200 and 'data-page="lineage"' in r.text
    nav = re.findall(r'<a href="(/[a-z]*)(?:\?scan=[^"]*)?" class="[^"]*">([^<]+)</a>', r.text.split("</nav>")[0])
    assert nav[-1] == ("/lineage", "Lineage") and nav[-2][1] == "Scans"
    for needle in ("Export CSV", "Export Excel", "Orphans", "Code hosted on", "Deploy target", "Environment tier", "Prod deployment", "GitHub org discovery is off", "display name"):
        assert needle in r.text
    assert 'hx-trigger="click once"' in r.text and "data-flow-toggle" in r.text  # server-rendered rows, the flow is loaded on first open
    assert "Traceback" not in r.text and "UndefinedError" not in r.text


def test_repo_page_fragment_and_links(client):
    repo = next(r for r in client.get("/api/v1/lineage").json()["repos"] if r["summary"]["releases"] and r["summary"]["pipelines"] and "/" in r["repo"]["name"])
    key = f"{repo['repo']['project']}/{repo['repo']['name']}"
    page = client.get(f"/lineage/{key}", params={"view": "table"})
    assert page.status_code == 200 and repo["repo"]["name"] in page.text and "Deployment" not in page.text.split("<h1")[0]
    for needle in ("Artifact source", "Last deployment", "Targets", "Service connections", "Approvals / gates", "Compliance"):
        assert needle in page.text
    frag = client.get(f"/lineage/{key}", params={"fragment": 1, "view": "table"})
    assert frag.status_code == 200 and "<html" not in frag.text and "Last deployment" in frag.text
    assert f"/lineage/{key}" in client.get("/lineage", params={"q": repo["repo"]["name"]}).text  # list links to the repo page
    assert f"/lineage/{key}" in client.get(f"/repos/{key}").text and f"/lineage/{key}" in client.get("/repos").text  # and so do the repo page and the table
    assert client.get(f"/api/v1/lineage/{key}").json()["repo"]["key"] == key
    assert client.get("/lineage/Payments/does-not-exist").status_code == 404
    assert client.get("/api/v1/lineage/Payments/does-not-exist").status_code == 404


def test_page_shows_chain_details(client):
    docs = client.get("/api/v1/lineage").json()["repos"]
    deploy = next(d for d in docs if any(s["last_deploy"]["status"] == "succeeded" and s["last_deploy"]["triggered_by"] for r in d["releases"] for s in r["stages"]))
    html = client.get(f"/lineage/{deploy['repo']['key']}", params={"view": "table"}).text
    ld = next(s["last_deploy"] for r in deploy["releases"] for s in r["stages"] if s["last_deploy"]["triggered_by"])
    assert ld["triggered_by"] in html and "@" not in ld["triggered_by"] and "succeeded" in html
    adopted = next(d for d in docs if any(p["adopted_from"] for p in d["pipelines"]))
    assert "defined in" in client.get(f"/lineage/{adopted['repo']['key']}", params={"view": "table"}).text
    assert "defined in" in client.get(f"/lineage/{adopted['repo']['key']}").text  # and in the flow
    other = next(d for d in docs for p in d["pipelines"] if p["yaml_in_other_repo"] and not p["adopted_from"])
    assert "YAML in another repo" in client.get(f"/lineage/{other['repo']['key']}", params={"view": "table"}).text
    consumer = next(d for d in docs for p in d["pipelines"] if p["downstream"])
    assert "Consumed by" in client.get(f"/lineage/{consumer['repo']['key']}", params={"view": "table"}).text
    assert "triggers pipeline" in client.get(f"/lineage/{consumer['repo']['key']}").text  # downstream = side branch in the flow


# ------------------------------------------------------------------ filters
def test_filters_on_page_and_api(client):
    allj = client.get("/api/v1/lineage").json()
    assert allj["summary"]["repos"] == len(allj["repos"]) > 100 and allj["has_data"]
    projects = allj["options"]["projects"]
    assert len(projects) >= 2
    p = client.get("/api/v1/lineage", params={"project": projects[0]}).json()["repos"]
    assert p and all(r["repo"]["project"] == projects[0] for r in p) and len(p) < len(allj["repos"])
    gh = client.get("/api/v1/lineage", params={"provider": "github"}).json()["repos"]
    assert gh and all(r["repo"]["provider"] == "github" for r in gh)
    aks = client.get("/api/v1/lineage", params={"target": "aks"}).json()["repos"]
    assert aks and all("aks" in r["summary"]["targets"] for r in aks)
    prod = client.get("/api/v1/lineage", params={"tier": "prod"}).json()["repos"]
    assert prod and all("prod" in r["summary"]["tiers"] for r in prod)
    yes = client.get("/api/v1/lineage", params={"has_prod": "yes"}).json()["repos"]
    no = client.get("/api/v1/lineage", params={"has_prod": "no"}).json()["repos"]
    assert yes and no and all(r["summary"]["has_prod"] for r in yes) and not any(r["summary"]["has_prod"] for r in no) and len(yes) + len(no) == len(allj["repos"])
    name = aks[0]["repo"]["name"]
    q = client.get("/api/v1/lineage", params={"q": name.upper()}).json()["repos"]
    assert any(r["repo"]["name"] == name for r in q) and len(q) < len(allj["repos"])
    res = next(t["name"] for r in aks for rel in r["releases"] for st in rel["stages"] for t in st["targets"] if t["name"])
    assert client.get("/api/v1/lineage", params={"q": res}).json()["repos"]  # search reaches resource names
    orph = client.get("/api/v1/lineage", params={"orphans": "1"}).json()
    assert orph["repos"] and all(r["summary"]["empty"] for r in orph["repos"])
    assert {o["type"] for o in orph["orphans"]} >= {"pipeline", "release", "repo"}
    assert client.get("/api/v1/lineage", params={"project": "No-Such-Project"}).json()["repos"] == []
    # <select> values arrive as empty strings: they mean "all", garbage is ignored, none of it is an error
    r = client.get("/lineage", params={"scan": "", "project": "", "provider": "", "target": "", "tier": "", "has_prod": "", "orphans": "", "q": ""})
    assert r.status_code == 200
    assert client.get("/lineage", params={"has_prod": "maybe"}).status_code == 200
    assert client.get("/lineage", params={"q": "x" * 5000}).status_code == 400
    page = client.get("/lineage", params={"orphans": "1", "project": projects[0]}).text
    assert "<details id=\"orphans\"" in page and " open>" in page.split('id="orphans"')[1][:200]


def test_orphans_section_lists_reasons(client):
    items = client.get("/api/v1/lineage").json()["orphans"]
    reasons = " | ".join(o["reason"] for o in items)
    for needle in ("no resolvable repository", "not found in this project", "no artifact source", "cannot be linked to a repository", "no GitHub reader is configured"):
        assert needle in reasons


# ------------------------------------------------------------------ CSV
def test_csv_columns_rows_and_headers(client):
    r = client.get("/lineage.csv")
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf") and r.headers["content-type"].startswith("text/csv")
    assert r.headers["cache-control"] == "no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["content-disposition"] == 'attachment; filename="pch-lineage-2026-10-01.csv"'
    rows = rows_of(r)
    assert rows[0] == LQ.HEADERS and len(set(LQ.HEADERS)) == len(LQ.HEADERS) and len(LQ.HEADERS) == 47
    assert all(len(x) == len(LQ.HEADERS) for x in rows)
    col = {h: i for i, h in enumerate(rows[0])}
    body = rows[1:]
    docs = client.get("/api/v1/lineage").json()["repos"]
    assert {x[col["repo"]] for x in body} == {d["repo"]["name"] for d in docs}  # every repo is present, incl. those without pipelines
    empty = next(d for d in docs if d["summary"]["empty"])
    only = [x for x in body if x[col["repo"]] == empty["repo"]["name"]]
    assert len(only) == 1 and only[0][col["pipeline_id"]] == "" and only[0][col["project"]] == empty["repo"]["project"]
    full = next(d for d in docs if d["releases"] and d["pipelines"] and len(d["releases"][0]["stages"]) >= 2)
    path_rows = [x for x in body if x[col["repo"]] == full["repo"]["name"] and x[col["release_kind"]] == "classic_release"]
    assert [x[col["stage"]] for x in path_rows[: len(full["releases"][0]["stages"])]] == [s["name"] for s in full["releases"][0]["stages"]]
    assert [x[col["stage_order"]] for x in path_rows[:2]] == ["1", "2"] and path_rows[0][col["release_name"]] == full["releases"][0]["name"]
    done = next(x for x in body if x[col["last_deploy_status"]] == "succeeded" and x[col["last_deploy_by"]])
    assert re.fullmatch(r"\d{4}-\d\d-\d\d \d\d:\d\d:\d\d", done[col["last_deploy_time_utc"]]) and "@" not in done[col["last_deploy_by"]] and done[col["last_deploy_version"]]
    assert any(x[col["yaml_in_other_repo"]] == "yes" for x in body) and any(x[col["downstream_pipelines"]] for x in body)
    assert any(x[col["target_resources"]] for x in body) and any(x[col["approvals_gates"]] for x in body)
    assert not any("@" in v for x in body for v in (x[col["last_deploy_by"]], x[col["ci_last_run_number"]]))


def test_exports_honour_filters_and_scan_selection(client):
    allrows = rows_of(client.get("/lineage.csv"))
    col = {h: i for i, h in enumerate(allrows[0])}
    project = allrows[1][col["project"]]
    f = rows_of(client.get("/lineage.csv", params={"project": project}))
    assert 1 < len(f) < len(allrows) and {x[col["project"]] for x in f[1:]} == {project}
    t = rows_of(client.get("/lineage.csv", params={"target": "sql", "tier": "prod"}))
    assert 1 < len(t) < len(allrows)
    for repo in {x[col["repo"]] for x in t[1:]}:  # every repo kept by the filter really deploys to SQL and has a prod stage
        mine = [x for x in t[1:] if x[col["repo"]] == repo]
        assert any("sql" in x[col["deploy_targets"]] for x in mine) and any(x[col["tier"]] == "prod" for x in mine), repo
    cd = client.get("/lineage.csv", params={"project": project, "scan": "scan-0"}).headers["content-disposition"]
    assert cd == f'attachment; filename="pch-lineage-{project}-2026-10-01.csv"'
    assert len(rows_of(client.get("/lineage.csv", params={"q": "zzz-nothing-matches"}))) == 1  # header only
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/lineage.xlsx", params={"project": project}).content))
    assert wb["Lineage"].max_row == len(f)  # the Excel export has exactly the rows the CSV has
    assert wb["Summary"]["B6"].value == f"project={project}"
    unknown = client.get("/lineage.csv", params={"scan": "does-not-exist"})  # unknown scan falls back to the latest, like every page
    assert unknown.status_code == 200


def test_csv_neutralises_formulas_in_every_cell(db, tmp_path):
    c = TestClient(create_app(_with_names(db, tmp_path, FORMULAS), settings=S()), base_url="http://localhost")
    text = c.get("/lineage.csv").content.decode("utf-8-sig")
    cells = [v for row in csv.reader(io.StringIO(text)) for v in row]
    assert any(v.startswith("'=HYPERLINK") for v in cells) and any(v.startswith("'+cmd") for v in cells) and any(v.startswith("'-2+3") for v in cells) and any(v.startswith("'@SUM") for v in cells)
    assert not [v for v in cells if v and v[0] in "=+-@"]


# ------------------------------------------------------------------ Excel
def test_xlsx_workbook_structure(client):
    r = client.get("/lineage.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == X.XLSX_MIME and r.headers["cache-control"] == "no-store"
    assert r.headers["content-disposition"] == 'attachment; filename="pch-lineage-2026-10-01.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Lineage", "Orphans"]
    ws = wb["Lineage"]
    assert [c.value for c in ws[1]] == LQ.HEADERS and ws.freeze_panes == "A2" and ws.auto_filter.ref == f"A1:AU{ws.max_row}"
    assert ws.column_dimensions["B"].width == 34 and ws.column_dimensions["AM"].width == 18
    csv_rows = rows_of(client.get("/lineage.csv"))
    assert ws.max_row == len(csv_rows)
    dates = [c for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, datetime)]
    assert dates and all(c.number_format == "yyyy-mm-dd hh:mm" for c in dates) and {ws.cell(1, c.column).value for c in dates} <= {"ci_last_run_time_utc", "last_deploy_time_utc"}
    assert all(c.data_type == "d" or c.data_type == "n" for c in dates)
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["Scan"] == "scan-0" and summary["Scan started (UTC)"] == datetime(2026, 10, 1, 12, 0, 0) and summary["Scan mode"] == "demo"
    assert summary["Rows in the Lineage sheet"] == len(csv_rows) - 1 and summary["Repositories"] >= 100
    labels = [row[0] for row in wb["Summary"].iter_rows(values_only=True) if row]
    for needle in ("Repositories by project", "Repositories by code host", "Repositories by deploy target", "Stages by environment tier", "Stages by last deployment status", "GitHub", "Payments", "prod"):
        assert needle in labels, needle
    orph = wb["Orphans"]
    assert [c.value for c in orph[1]] == ["type", "project", "name", "id", "url", "reason"] and orph.max_row > 5
    assert {orph.cell(i, 1).value for i in range(2, orph.max_row + 1)} >= {"pipeline", "release", "repo"}


def test_xlsx_never_contains_formulas(db, tmp_path):
    """Regression: attacker-controlled names that start with = + - @ must be literal text in every sheet."""
    c = TestClient(create_app(_with_names(db, tmp_path, FORMULAS), settings=S()), base_url="http://localhost")
    wb = openpyxl.load_workbook(io.BytesIO(c.get("/lineage.xlsx").content))
    seen = 0
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith("="), (ws.title, cell.coordinate)
                    seen += "HYPERLINK" in cell.value or "SUM(1+1)" in cell.value or "cmd|" in cell.value
    assert seen >= 4  # the hostile names are present (as text), not dropped
    with __import__("zipfile").ZipFile(io.BytesIO(c.get("/lineage.xlsx").content)) as z:  # and no <f> formula element in any sheet XML
        assert not any(b"<f>" in z.read(n) or b"<f " in z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))


def test_string_cell_is_literal_even_without_neutralisation(monkeypatch):
    """The explicit string data type alone protects against formulas (defence in depth behind csv_cell)."""
    monkeypatch.setattr(X, "csv_cell", lambda v: v)
    cols = [("a", "a", 10)]
    data = X.build_xlsx(cols, [["=1+1"], ["+1+1"], ["-1+1"], ["@SUM(1)"], ["=cmd|' /C calc'!A0"]], [["x", "=1+1"]], cols, [["=2+2"]])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate)
    assert wb["Lineage"]["A2"].value == "=1+1" and wb["Lineage"]["A2"].data_type == "s"
    assert wb["Summary"]["B1"].data_type == "s" and wb["Orphans"]["A2"].data_type == "s"


def test_illegal_control_characters_do_not_break_the_workbook():
    cols = [("a", "a", 10)]
    wb = openpyxl.load_workbook(io.BytesIO(X.build_xlsx(cols, [["bad\x00\x07name\x1f"]], [], cols, [])))
    assert wb["Lineage"]["A2"].value == "badname"


# ------------------------------------------------------------------ row cap
def test_export_row_cap_is_enforced_with_a_clear_message(db):
    c = TestClient(create_app(db, settings=S(export_max_rows=5)), base_url="http://localhost")
    for path in ("/lineage.csv", "/lineage.xlsx"):
        r = c.get(path)
        assert r.status_code == 413 and "more than 5 rows" in r.text and "EXPORT_MAX_ROWS" in r.text and "Narrow the filters" in r.text, path
        assert r.headers["cache-control"] == "no-store" and "Traceback" not in r.text
    ok = c.get("/lineage.csv", params={"q": "zzz-nothing-matches"})
    assert ok.status_code == 200
    one = rows_of(TestClient(create_app(db, settings=S(export_max_rows=1_000_000)), base_url="http://localhost").get("/lineage.csv"))
    n = len(one) - 1
    edge = TestClient(create_app(db, settings=S(export_max_rows=n)), base_url="http://localhost")
    assert edge.get("/lineage.csv").status_code == 200  # exactly at the limit is allowed
    assert TestClient(create_app(db, settings=S(export_max_rows=n - 1)), base_url="http://localhost").get("/lineage.csv").status_code == 413
    j = c.get("/api/v1/lineage")  # the JSON API is not an export: not capped
    assert j.status_code == 200


# ------------------------------------------------------------------ method guard + auth
def test_export_and_pages_are_get_only(client):
    for path in ("/lineage", "/lineage.csv", "/lineage.xlsx", "/api/v1/lineage", "/lineage/Payments/x"):
        for method in ("post", "put", "delete", "patch"):
            r = getattr(client, method)(path)
            assert r.status_code == 405, (method, path, r.status_code)


def test_lineage_requires_authentication_and_role(db):
    def hdr(roles=(ROLE,)):
        claims = [{"typ": "http://schemas.microsoft.com/identity/claims/objectidentifier", "val": "oid-1"}, {"typ": "name", "val": "Ada"}] + [{"typ": "roles", "val": r} for r in roles]
        return {"X-MS-CLIENT-PRINCIPAL": base64.b64encode(json.dumps({"auth_typ": "aad", "name_typ": "name", "role_typ": "roles", "claims": claims}).encode()).decode()}

    c = TestClient(create_app(db, settings=S(app_env="dev", auth_mode="easyauth", auth_allowed_roles=ROLE, auth_easyauth_assume_enabled=True)), base_url="http://localhost")
    paths = ("/lineage", "/lineage.csv", "/lineage.xlsx", "/api/v1/lineage", "/lineage/Payments/x", "/api/v1/lineage/Payments/x")
    for path in paths:
        r = c.get(path)
        assert r.status_code == 401, path
        assert "scan-0" not in r.text and "PK" not in r.text[:2] and "project," not in r.text
        r = c.get(path, headers=hdr(roles=("Other",)))
        assert r.status_code == 403 and "<nav" not in r.text, path
    for path in ("/lineage", "/lineage.csv", "/lineage.xlsx", "/api/v1/lineage"):
        assert c.get(path, headers=hdr()).status_code == 200, path


# ------------------------------------------------------------------ escaping
def _with_names(db: str, tmp: Path, names: list[str]) -> str:
    """A copy of the scan DB whose lineage documents carry hostile names in every text field the UI and exports show."""
    dest = tmp / "hostile.db"
    copy_db(db, dest)
    url = f"sqlite:///{dest}"
    with session_scope(url) as s:
        rows = [r for r in s.query(LineageRow).filter(LineageRow.kind == "repo").all() if r.n_pipelines and r.n_releases][:4]
        assert len(rows) == 4
        for r, p in zip(rows, names, strict=True):
            d = json.loads(json.dumps(r.doc))
            d["repo"]["name"] = p
            d["repo"]["key"] = f"{d['repo']['project']}/{p}"
            d["repo"]["url"] = p
            d["repo"]["service_connection"] = p
            r.repo_key = d["repo"]["key"]
            for pl in d["pipelines"]:
                pl.update(name=p, url=p, definition_path=p, yaml_repo=p, artifacts=[p], code_repos=[p])
                pl["trigger"]["ci_branches"] = [p]
                pl["upstream"] = [{"kind": "yaml_resource", "name": p, "detail": p, "url": p}]
                for st in pl["stages"]:
                    st.update(name=p, env_name=p)
            for rel in d["releases"]:
                rel.update(name=p, url=p)
                rel["sources"] = [{"type": p, "alias": p, "name": p, "primary": True}]
                for st in rel["stages"]:
                    st.update(name=p, env_name=p, service_connections=[p], approvals=[p], gates=[p], depends_on=[p])
                    st["targets"] = [{"kind": "webapp", "name": p, "detail": p}]
                    st["last_deploy"].update(version=p, artifact_version=p, triggered_by=p, url=p, note=p, status="succeeded")
            r.doc = d
        s.add(LineageRow(scan_id="scan-0", repo_key=f"{rows[0].project}/(unlinked)", kind="orphan", project=rows[0].project, doc={"orphans": [
            {"type": "pipeline", "project": rows[0].project, "id": p, "name": p, "url": p, "reason": p} for p in names]}))
    return url


def test_xss_escaped_everywhere_on_lineage_pages(db, tmp_path):
    url = _with_names(db, tmp_path, PAYLOADS)
    c = TestClient(create_app(url, settings=S()), base_url="http://localhost")
    pages = ["/lineage", "/lineage?orphans=1", "/lineage?q=img"]
    docs = c.get("/api/v1/lineage").json()["repos"]
    hostile = [d for d in docs if d["repo"]["name"] in PAYLOADS]
    assert len(hostile) == 4
    for d in hostile:
        from urllib.parse import quote

        base = f"/lineage/{quote(d['repo']['project'], safe='')}/{quote(d['repo']['name'], safe='/')}"
        pages += [base, base + "?fragment=1"]
    escaped = 0
    for path in pages:
        r = c.get(path)
        assert r.status_code == 200, path
        t = r.text
        for p in ("<img src=x onerror=alert(1)>", "</script><script>alert(1)", '"><svg/onload', "<svg/onload=alert(1)>"):
            assert p not in t, (path, p)
        assert 'href="javascript:' not in t
        escaped += "&lt;img src=x onerror=alert(1)&gt;" in t
        assert not re.findall(r"<script(?![^>]*\bsrc=)(?![^>]*application/json)[^>]*>", t)
    assert escaped >= 3
    api = c.get("/api/v1/lineage")
    assert api.headers["content-type"].startswith("application/json") and api.headers["x-content-type-options"] == "nosniff"


def test_pre_lineage_scan_degrades_gracefully(db, tmp_path):
    dest = tmp_path / "old.db"
    copy_db(db, dest)
    url = f"sqlite:///{dest}"
    with session_scope(url) as s:
        s.query(LineageRow).delete()
    c = TestClient(create_app(url, settings=S()), base_url="http://localhost")
    r = c.get("/lineage")
    assert r.status_code == 200 and "no lineage data" in r.text and "Export CSV" in r.text
    for path in ("/lineage.csv", "/lineage.xlsx"):
        e = c.get(path)
        assert e.status_code == 404 and "no lineage data" in e.text
    assert c.get("/api/v1/lineage").json()["has_data"] is False


# ------------------------------------------------------------------ docs
def test_readme_documents_every_export_column_in_order():
    readme = (Path(__file__).parent.parent / "README.md").read_text()
    section = readme.split("**Export columns**", 1)[1].split("\n## ", 1)[0]
    documented = re.findall(r"^\| (\d+) \| `([a-z_]+)` \|", section, re.M)
    assert [int(n) for n, _ in documented] == list(range(1, len(LQ.HEADERS) + 1))
    assert [c for _, c in documented] == LQ.HEADERS
    assert {k for k, _, _ in LQ.COLUMNS if k.endswith("_time_utc")} == LQ.DATE_COLUMNS
