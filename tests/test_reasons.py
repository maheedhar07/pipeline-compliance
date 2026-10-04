"""Reasons for non-compliance (G1): computation, persistence, and every place they are shown or exported."""

import asyncio
import csv
import io
from datetime import datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from pch.demo.generator import generate_world
from pch.demo.transport import parse_world_time
from pch.engine.reasons import compute_reasons, short_title
from pch.model.findings import Finding, Severity, Status
from pch.orchestrator import ScanConfig, Scanner
from pch.settings import Policy, Scope
from pch.sources import demo_sources
from pch.store.db import session_scope
from pch.store.models import RepoResultRow
from pch.web.app import create_app
from pch.web.queries import top_reasons


def F(rule_id, sev, status, msg="", pipeline=None, stage=None, category=None):
    return Finding(rule_id=rule_id, repo_key="P/r", category=category or rule_id[:3], severity=sev, status=status, message=msg, pipeline_id="1" if pipeline else None,
                   pipeline_name=pipeline, stage=stage)


# ------------------------------------------------------------------ computation
def test_short_title_drops_parentheses_and_truncates():
    assert short_title("Default branch requires 2+ reviewers (creator vote excluded, reset on push)") == "Default branch requires 2+ reviewers"
    long = short_title("x" * 200)
    assert len(long) == 64 and long.endswith("…")


def test_reasons_are_failing_rules_ordered_by_severity_and_deduped_per_rule():
    fs = [
        F("QLT-004", Severity.MEDIUM, Status.FAIL, "uses 'Sonar way'"),
        F("DEP-001", Severity.CRITICAL, Status.FAIL, "no manual approval", "rel", "Prod"),
        F("DEP-001", Severity.CRITICAL, Status.FAIL, "no manual approval b", "rel2", "Prod"),
        F("DEP-001", Severity.CRITICAL, Status.PASS, "ok", "rel3", "Prod"),
        F("SRC-004", Severity.HIGH, Status.FAIL, "classic definition"),
        F("TST-001", Severity.HIGH, Status.PASS),
        F("HYG-003", Severity.INFO, Status.WARN, "w"),
        F("SEC-005", Severity.MEDIUM, Status.WARN),
        F("SRC-001", Severity.HIGH, Status.UNKNOWN),
        F("SUP-005", Severity.LOW, Status.WAIVED),
        F("COLLECTION-ERROR", Severity.INFO, Status.UNKNOWN, "ado: boom", category="SYS"),
    ]
    r = compute_reasons(fs)
    assert [i["rule_id"] for i in r["items"]] == ["DEP-001", "SRC-004", "QLT-004"]  # critical, high, medium; one per rule
    dep = r["items"][0]
    assert dep["count"] == 2 and dep["severity"] == "critical" and dep["message"] == "no manual approval"  # first by pipeline name
    assert dep["label"] == "DEP-001 Production deployments require manual approval"
    assert dep["text"] == "DEP-001 Production deployments require manual approval: no manual approval"
    assert (r["fail"], r["warn"], r["unknown"]) == (3, 2, 1)  # waived and SYS are not counted anywhere


def test_a_passing_or_empty_repo_has_no_reasons():
    assert compute_reasons([]) == {"items": [], "fail": 0, "warn": 0, "unknown": 0}
    assert compute_reasons([F("SRC-004", Severity.HIGH, Status.PASS)])["items"] == []


def test_message_is_collapsed_and_clipped():
    r = compute_reasons([F("SRC-004", Severity.HIGH, Status.FAIL, "a\n  b " + "z" * 500)])
    assert "\n" not in r["items"][0]["message"] and len(r["items"][0]["message"]) <= 220 and r["items"][0]["message"].endswith("…")


def test_top_reasons_counts_repos():
    rows = [{"reasons": [{"rule_id": "A-1", "label": "A-1 a", "title": "a", "severity": "high"}, {"rule_id": "B-1", "label": "B-1 b", "title": "b", "severity": "critical"}]},
            {"reasons": [{"rule_id": "A-1", "label": "A-1 a", "title": "a", "severity": "high"}]}, {"reasons": []}]
    assert [(t["rule_id"], t["repos"]) for t in top_reasons(rows)] == [("A-1", 2), ("B-1", 1)]


# ------------------------------------------------------------------ persisted + pages + exports
@pytest.fixture(scope="module")
def env(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("reasons")
    db = f"sqlite:///{tmp}/r.db"
    w = generate_world(seed=11, repos=70, now=datetime(2026, 10, 1, 12, 0, 0))
    src = demo_sources(w)
    cfg = ScanConfig(scope=Scope(projects=w["meta"]["projects"]), policy=Policy(approved_registries=["contosoacr.azurecr.io"]), db_url=db, mode="demo", now=parse_world_time(w))

    async def go():
        try:
            await Scanner(src, cfg).run("rs")
        finally:
            await src.aclose()

    asyncio.run(go())
    return TestClient(create_app(db)), db


def rows_of(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


def test_reasons_are_persisted_per_repo(env):
    _, db = env
    with session_scope(db) as s:
        rows = list(s.query(RepoResultRow).filter(RepoResultRow.scan_id == "rs"))
    assert rows
    bad = [r for r in rows if r.status in ("NON_COMPLIANT", "AT_RISK")]
    assert bad and any(r.external["reasons"]["items"] for r in bad)
    for r in rows:
        rs = r.external["reasons"]
        assert set(rs) == {"items", "fail", "warn", "unknown"} and rs["fail"] == len(rs["items"])
        sevs = [i["severity"] for i in rs["items"]]
        order = ["critical", "high", "medium", "low", "info"]
        assert sevs == sorted(sevs, key=order.index)
        if r.status == "NON_COMPLIANT":
            assert rs["items"][0]["severity"] == "critical"  # a critical failure is what makes a repo non-compliant


def test_repos_api_csv_and_page_show_why(env):
    c, _ = env
    repos = c.get("/api/v1/repos").json()["repos"]
    nc = [r for r in repos if r["status"] == "NON_COMPLIANT"]
    assert nc and all(r["reasons"] and r["why"] == r["reasons"][:3] and r["why_more"] == max(0, len(r["reasons"]) - 3) for r in nc)
    html = c.get("/repos").text
    assert ">Why</th>" in html
    sample = next(r for r in nc if r["why_more"])
    assert sample["why"][0]["label"].split(" ", 1)[0] in html and f"+{sample['why_more']} more" in html
    comp = next(r for r in repos if r["status"] == "COMPLIANT")
    assert f'/repos/{comp["project"]}/{comp["repo"]}' in html  # compliant repos stay listed without reasons
    rows = rows_of(c.get("/repos.csv"))
    assert rows[0][-1] == "reasons" and {"score", "status"} <= set(rows[0])
    by_repo = {(r[0], r[1]): r for r in rows[1:]}
    got = by_repo[(sample["project"], sample["repo"])][-1]
    assert got.split("; ")[0] == sample["reasons"][0]["text"] and len(got.split("; ")) >= len(sample["reasons"])


def test_repos_filter_by_failing_rule(env):
    c, _ = env
    rid = c.get("/api/v1/overview").json()["top_reasons"][0]["rule_id"]
    hit = c.get("/api/v1/repos", params={"rule": rid}).json()["repos"]
    assert hit and all(any(i["rule_id"] == rid for i in r["reasons"]) for r in hit)
    assert len(hit) < len(c.get("/api/v1/repos").json()["repos"])
    page = c.get("/repos", params={"rule": rid}).text
    assert "fails." in page and f"/rules/{rid}" in page


def test_overview_has_top_reasons_and_noncompliant_list(env):
    c, _ = env
    ov = c.get("/api/v1/overview").json()
    top = ov["top_reasons"]
    assert top and [t["repos"] for t in top] == sorted((t["repos"] for t in top), reverse=True)
    nc = ov["noncompliant"]
    assert nc and nc == sorted(nc, key=lambda r: (r["score"] if r["score"] is not None else 101, r["project"], r["repo"])) and all(r["why"] for r in nc)
    html = c.get("/").text
    assert "Top reasons" in html and "Non-compliant repositories" in html
    assert f"/findings?rule={top[0]['rule_id']}&amp;status=FAIL" in html and f"/repos?rule={top[0]['rule_id']}" in html
    assert f'/repos/{nc[0]["project"]}/{nc[0]["repo"]}' in html


def test_repo_page_has_why_box(env):
    c, _ = env
    r = next(r for r in c.get("/api/v1/repos").json()["repos"] if r["status"] == "NON_COMPLIANT")
    html = c.get(f"/repos/{r['project']}/{r['repo']}").text
    assert "Why this repo is non compliant" in html
    assert html.index("Why this repo is non compliant") < html.index("Test state")  # at the top
    assert r["reasons"][0]["rule_id"] in html and r["reasons"][0]["title"] in html
    ok = next((x for x in c.get("/api/v1/repos").json()["repos"] if x["status"] == "COMPLIANT" and not x["reasons"]), None)
    if ok:
        assert "Why this repo is" not in c.get(f"/repos/{ok['project']}/{ok['repo']}").text


def test_lineage_rows_show_status_reasons_and_failing_rules(env):
    c, _ = env
    d = c.get("/api/v1/lineage").json()
    with_reasons = [x for x in d["repos"] if x["compliance"] and x["compliance"]["reasons"]]
    assert with_reasons and all(x["compliance"]["status"] for x in d["repos"] if x["compliance"])
    html = c.get("/lineage").text
    assert "why: " in html
    repo = next(x for x in with_reasons if not x["summary"]["empty"])
    detail = c.get(f"/api/v1/lineage/{repo['repo']['project']}/{repo['repo']['name']}").json()
    fails = [p for p in [*detail["pipelines"], *detail["releases"]] if p["failing"] or any(st["failing"] for st in p["stages"])]
    assert all("failing" in p and all("failing" in st for st in p["stages"]) for p in [*detail["pipelines"], *detail["releases"]])
    assert fails  # pipeline/stage level findings (e.g. SRC-004, DEP-*) are attached
    f = next(x for p in fails for x in (p["failing"] or [y for st in p["stages"] for y in st["failing"]]))
    assert f["rule_id"] and f["severity"] in ("critical", "high", "medium", "low", "info")
    frag = c.get(f"/lineage/{repo['repo']['project']}/{repo['repo']['name']}", params={"fragment": "1"}).text
    assert "failing:" in frag and f["rule_id"] in frag
    page = c.get(f"/lineage/{repo['repo']['project']}/{repo['repo']['name']}").text
    assert "Compliance of this repository" in page


def test_lineage_csv_and_xlsx_carry_status_score_reasons(env):
    c, _ = env
    rows = rows_of(c.get("/lineage.csv"))
    assert rows[0][-3:] == ["status", "score", "reasons"]
    col = {h: i for i, h in enumerate(rows[0])}
    repos = {r["repo"]: r for r in c.get("/api/v1/repos").json()["repos"]}
    seen = 0
    for r in rows[1:]:
        api = next(v for k, v in repos.items() if k == r[col["repo"]] or k.endswith("/" + r[col["repo"]]) or v["repo"] == r[col["repo"]])
        assert r[col["status"]] == api["status"]
        assert r[col["reasons"]] == "; ".join(i["text"] for i in api["reasons"])
        assert r[col["score"]] == ("" if api["score"] is None else str(api["score"]))
        seen += bool(r[col["reasons"]])
    assert seen
    wb = load_workbook(BytesIO(c.get("/lineage.xlsx").content))
    summary = [[cell.value for cell in row] for row in wb["Summary"].iter_rows()]
    flat = [str(v) for row in summary for v in row if v is not None]
    assert "Top reasons (failing rules, repositories affected)" in flat and "Repositories by compliance status" in flat
    head = [cell.value for cell in wb["Lineage"][1]]
    assert head[-3:] == ["status", "score", "reasons"]


def test_reasons_are_escaped_everywhere(env):
    c, db = env
    evil = "<script>alert(1)</script>"
    with session_scope(db) as s:
        row = s.query(RepoResultRow).filter(RepoResultRow.scan_id == "rs", RepoResultRow.status == "NON_COMPLIANT").first()
        ext = dict(row.external)
        rs = dict(ext["reasons"])
        rs["items"] = [{**rs["items"][0], "message": evil, "text": f"X {evil}", "label": f"X {evil}", "title": evil}, *rs["items"][1:]]
        ext["reasons"] = rs
        row.external = ext
        project, repo = row.project, row.repo
    for path in ("/", "/repos", f"/repos/{project}/{repo}", "/lineage", f"/lineage/{project}/{repo}"):
        html = c.get(path).text
        assert evil not in html, path
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in c.get(f"/repos/{project}/{repo}").text
    assert "<script>alert" not in c.get("/").text
