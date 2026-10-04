"""CSV and Excel (.xlsx) writers for tabular exports.

Spreadsheet formula injection: repo, pipeline, stage and environment names are attacker-influenced. Every string is
(1) neutralised with a leading apostrophe when it starts with ``= + - @`` / control whitespace (``csv_cell``, the same rule
as ``/repos.csv``) and (2) in Excel written with an explicit ``string`` cell type, so openpyxl can never turn a value into
a formula (it would for any str starting with "=").
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Iterator, Sequence
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BOM = "﻿"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # control characters XML (and openpyxl) cannot hold
MAX_CELL = 32000  # Excel limit is 32767 characters


def csv_cell(v: Any) -> Any:
    """Neutralise spreadsheet formula injection (repo names etc. are attacker-influenced)."""
    if isinstance(v, str) and v and (v[0] in "=+-@\t\r" or v.startswith(("\n",))):
        return "'" + v
    return v


def clean_text(v: str) -> str:
    return _ILLEGAL.sub("", v)[:MAX_CELL]


def _csv_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, str):
        return csv_cell(clean_text(v))
    if v is None:
        return ""
    return v


def csv_lines(headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> Iterator[str]:
    """UTF-8 text lines (CRLF) starting with a BOM so Excel detects the encoding. Stable column order = ``headers``."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(list(headers))
    yield BOM + buf.getvalue()
    for r in rows:
        buf.seek(0)
        buf.truncate()
        w.writerow([_csv_value(c) for c in r])
        yield buf.getvalue()


# --------------------------------------------------------------------------- xlsx
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="3730A3")
DATE_FORMAT = "yyyy-mm-dd hh:mm"


def string_cell(ws: Any, value: str, *, bold: bool = False) -> WriteOnlyCell:
    """A text cell that is ALWAYS a literal string, whatever it starts with (explicit data_type, never a formula)."""
    c = WriteOnlyCell(ws, value=csv_cell(clean_text(value)))
    c.data_type = "s"
    if bold:
        c.font = Font(bold=True)
    return c


def _cell(ws: Any, v: Any) -> Any:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        c = WriteOnlyCell(ws, value=v)
        c.number_format = DATE_FORMAT
        return c
    if isinstance(v, bool):
        return string_cell(ws, "yes" if v else "no")
    if isinstance(v, int | float):
        return WriteOnlyCell(ws, value=v)
    return string_cell(ws, str(v))


def _header(ws: Any, headers: Sequence[str]) -> list[WriteOnlyCell]:
    out = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.data_type = "s"
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center")
        out.append(c)
    return out


def _sheet(wb: Workbook, title: str, columns: Sequence[tuple[str, str, int]]) -> Any:
    ws = wb.create_sheet(title)
    for i, (_, _, width) in enumerate(columns, 1):  # widths and panes must be set before the first row in write-only mode
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.append(_header(ws, [h for _, h, _ in columns]))
    return ws


def build_xlsx(
    columns: Sequence[tuple[str, str, int]], rows: Iterable[Sequence[Any]], summary_rows: Sequence[Sequence[Any]],
    orphan_columns: Sequence[tuple[str, str, int]], orphans: Iterable[Sequence[Any]],
    sheet_names: tuple[str, str, str] = ("Summary", "Lineage", "Orphans"),
) -> bytes:
    """Write-only (streaming) workbook: Summary, the data sheet (frozen header, autofilter, widths, real dates) and Orphans."""
    wb = Workbook(write_only=True)
    ws_sum = wb.create_sheet(sheet_names[0])
    ws_sum.column_dimensions["A"].width = 34
    ws_sum.column_dimensions["B"].width = 46
    for r in summary_rows:
        if r and r[0] == "#":  # section heading
            ws_sum.append([string_cell(ws_sum, str(r[1]), bold=True)])
        else:
            ws_sum.append([_cell(ws_sum, x) for x in r])
    ws = _sheet(wb, sheet_names[1], columns)
    n = 0
    for r in rows:
        ws.append([_cell(ws, v) for v in r])
        n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n + 1}"
    wo = _sheet(wb, sheet_names[2], orphan_columns)
    m = 0
    for r in orphans:
        wo.append([_cell(wo, v) for v in r])
        m += 1
    wo.auto_filter.ref = f"A1:{get_column_letter(len(orphan_columns))}{m + 1}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
